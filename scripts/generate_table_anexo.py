import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml

doc = docx.Document()
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(6)
p_title.paragraph_format.space_after = Pt(4)
run_title = p_title.add_run('Tabla Anexo C. Gobernanza del dato, métodos de ingesta y validación de calidad en dbt Core')
run_title.font.name = 'Verdana'
run_title.font.size = Pt(11)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(16, 44, 87)

table_data = [
    ['Capa / Dominio', 'Método de Ingesta (Bronze / Origen)', 'Archivo de Especificación', 'Modelos Cubiertos', 'Pruebas de Calidad (dbt)'],
    ['Gold', 'Transformación y materialización analítica dbt SQL (dbt run --select gold.*) a partir de Silver + enriquecimiento ML', 'models/gold/schema.yml', '8 modelos analíticos', '64 pruebas (unique, not_null, relationships, accepted_values)'],
    ['Silver Espacial', 'API CKAN (datos.tenerife.es), WFS IDECanarias/GRAFCAN, Overpass API (OSM) y exportación GEE → Azure Blob → PostGIS', 'models/silver/espacial/schema.yml', '9 modelos espaciales', '52 pruebas (integridad geométrica PostGIS, unique, not_null, relationships)'],
    ['Silver Clima', 'API REST v2.0.0 Agrocabildo con control de flujo (10 req/min) y partición Hive (año=YYYY/mes=MM/) → Azure Blob → PostgreSQL', 'models/silver/clima/schema.yml', '2 modelos climáticos', '15 pruebas (validación de sensores, integridad referencial de 1.8M filas)'],
    ['Silver Movilidad', 'Descarga directa de feeds ZIP GTFS (TITSA/Tranvía) y descarga automatizada de CSVs de AENA → Azure Blob → PostgreSQL', 'models/silver/movilidad/schema.yml', '3 modelos GTFS / AENA', '19 pruebas (stop_id, shape_id, aeropuertos canónicos TFN/TFS)'],
    ['Silver ISTAC', 'API REST / SDMX estadística del ISTAC (sistemas C00067A y C00065A_000061) con paginación JSON → Azure Blob → PostgreSQL', 'models/silver/istac/schema.yml', '3 modelos demográficos/laborales', '19 pruebas (códigos INE municipales 38001-38052, temporalidad)'],
    ['Silver Alojamiento', 'Extracción web / API abierta del Registro General Turístico (Gobierno de Canarias) + Geocodificación por lotes con caché PostGIS', 'models/silver/alojamiento/schema.yml', '1 modelo oficial', '7 pruebas (categorías oficiales, tipologías regladas, coordenadas)'],
    ['Silver Booking', 'Web scraping ético distribuido con rotación de User-Agent, delays probabilísticos (2,5–5 s) y sitemaps XML → Azure Blob → PostgreSQL', 'models/silver/booking/schema.yml', '2 modelos OTA Booking', '13 pruebas (integridad referencial review-hotel, ratings válidos)'],
    ['Silver TripAdvisor', 'Scraping / Terra API con validación geoespacial perimetral en PostGIS para evitar homónimos → Azure Blob → PostgreSQL', 'models/silver/tripadvisor/schema.yml', '2 modelos OTA TripAdvisor', '14 pruebas (location_id, reviews limpias > 15 caracteres)'],
    ['Silver YouTube', 'Extracción automatizada vía YouTube Data API v3 (Google Cloud Platform) con endpoints REST paginados → Azure Blob → PostgreSQL', 'models/silver/youtube/schema.yml', '2 modelos social video', '12 pruebas (video_id, comentarios válidos, no nulos)'],
    ['Silver LosViajeros', 'Web scraping ético en Python (BeautifulSoup) sobre paginación HTML de foros (248 hilos, cortesía 1,5 s) → Azure Blob → PostgreSQL', 'models/silver/losviajeros/schema.yml', '1 modelo foros de viajes', '5 pruebas (unique, not_null en mensaje, tema, texto y fecha)'],
    ['Singular Tests', 'Scripts Python y queries SQL de validación cruzada y privacidad', 'tests/', 'Pruebas transversales', 'assert_no_personal_data_columns (RGPD), assert_rating_in_range, etc.']
]

table = doc.add_table(rows=len(table_data), cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
col_widths = [Inches(1.1), Inches(2.2), Inches(1.5), Inches(1.1), Inches(1.8)]

for r_idx, row in enumerate(table.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.width = col_widths[c_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(table_data[r_idx][c_idx])
        run.font.name = 'Verdana'
        if r_idx == 0:
            shd = parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="102C57"/>')
            tcPr.append(shd)
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
        else:
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(40, 40, 40)
            if r_idx % 2 == 1:
                shd = parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F4F6F9"/>')
                tcPr.append(shd)

doc.save('docs/tabla_anexo_gobernanza_dbt.docx')
print('Successfully saved to docs/tabla_anexo_gobernanza_dbt.docx')
