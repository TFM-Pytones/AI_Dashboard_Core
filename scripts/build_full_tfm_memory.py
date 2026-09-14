"""
build_full_tfm_memory.py
========================
Compila la Memoria Oficial y Definitiva del Trabajo de Fin de Máster (UCM)
en formato Markdown (docs/TFM_Memoria_Completa.md y en el directorio de artefactos)
y en formato Microsoft Word (.docx) estilizado con tipografía Verdana,
paleta corporativa UCM (#1F4E79) y maquetación académica profesional.
"""

import os
import re
import sys
from pathlib import Path

# Añadir el directorio scripts al path
CURRENT_DIR = Path(__file__).resolve().parent
BASE_DIR = CURRENT_DIR.parent
sys.path.insert(0, str(CURRENT_DIR))

from memoria_sections_part1 import (
    get_front_matter,
    get_chapter_1,
    get_chapter_2,
    get_chapter_3,
)
from memoria_sections_part2 import (
    get_chapter_4,
    get_chapter_5,
)
from memoria_sections_part3 import (
    get_chapter_6,
    get_chapter_7,
    get_chapter_8,
    get_references,
    get_appendices,
)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# ==============================================================================
# 1. ENSAMBLADO DEL DOCUMENTO MARKDOWN MAESTRO
# ==============================================================================
def assemble_markdown():
    print("[1/4] Ensamblando secciones de la Memoria en formato Markdown...")
    sections = [
        get_front_matter().strip(),
        "\n\n---\n\n",
        get_chapter_1().strip(),
        "\n\n---\n\n",
        get_chapter_2().strip(),
        "\n\n---\n\n",
        get_chapter_3().strip(),
        "\n\n---\n\n",
        get_chapter_4().strip(),
        "\n\n---\n\n",
        get_chapter_5().strip(),
        "\n\n---\n\n",
        get_chapter_6().strip(),
        "\n\n---\n\n",
        get_chapter_7().strip(),
        "\n\n---\n\n",
        get_chapter_8().strip(),
        "\n\n---\n\n",
        get_references().strip(),
        "\n\n---\n\n",
        get_appendices().strip(),
    ]
    full_markdown = "".join(sections)

    # 1. Guardar en docs/TFM_Memoria_Completa.md
    docs_output = BASE_DIR / "docs" / "TFM_Memoria_Completa.md"
    docs_output.parent.mkdir(parents=True, exist_ok=True)
    with open(docs_output, "w", encoding="utf-8") as f:
        f.write(full_markdown)
    print(f" -> Guardado exitosamente en: {docs_output} ({len(full_markdown):,} caracteres)")

    # 2. Guardar en Artifacts Directory para inspección directa en el IDE
    artifact_dir = Path("C:/Users/ROBERTO/.gemini/antigravity-ide/brain/f560c0ab-af7d-45d0-9cd9-b885416bd59e")
    if artifact_dir.exists():
        artifact_output = artifact_dir / "TFM_Memoria_Completa.md"
        with open(artifact_output, "w", encoding="utf-8") as f:
            f.write(full_markdown)
        print(f" -> Guardado exitosamente en el directorio de artefactos: {artifact_output}")

    return full_markdown


# ==============================================================================
# 2. GENERACIÓN DEL DOCUMENTO WORD (.DOCX) CON PYTHON-DOCX
# ==============================================================================
def clean_xml_string(text: str) -> str:
    if not text:
        return ""
    # Filtrar caracteres de control no permitidos en XML 1.0
    return re.sub(r'[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD]', '', str(text))

def clean_math_and_latex(text: str) -> str:
    if not text:
        return ""
    t = str(text)
    # 1. Reemplazar fracciones \frac{a}{b} -> a / b
    t = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'\1 / \2', t)
    # 2. Reemplazar \sqrt{...} -> √(...)
    t = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', t)
    # 3. Reemplazar comandos de formato textual en LaTeX
    t = re.sub(r'\\text\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\textbf\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\operatorname\{([^}]+)\}', r'\1', t)
    t = re.sub(r'\\pmod\{([^}]+)\}', r'mod \1', t)
    t = re.sub(r'\\left|\\right', '', t)
    # 4. Reemplazo de símbolos LaTeX por caracteres Unicode limpios
    replacements = [
        (r'\le', '≤'),
        (r'\ge', '≥'),
        (r'\times', '×'),
        (r'\cdot', '·'),
        (r'\approx', '≈'),
        (r'\pm', '±'),
        (r'\,', ' '),
        (r'\%', '%'),
        (r'\circ', '°'),
        (r'\degree', '°'),
        (r'^\circ', '°'),
        (r'\partial', '∂'),
        (r'\sum', 'Σ'),
        (r'\alpha', 'α'),
        (r'\beta', 'β'),
        (r'\gamma', 'γ'),
        (r'\theta', 'θ'),
        (r'\pi', 'π'),
        (r'\sigma', 'σ'),
        (r'\mu', 'μ'),
        (r'\varepsilon', 'ε'),
        (r'\in', '∈'),
        (r'\infty', '∞'),
        (r'\_', '_'),
        (r'\quad', ' '),
        (r'\qquad', ' ')
    ]
    for pat, rep in replacements:
        t = t.replace(pat, rep)
    # 5. Eliminar delimitadores $ y $$
    t = re.sub(r'\$\$?', '', t)
    # 6. Limpiar espacios dobles
    t = re.sub(r' +', ' ', t)
    return t.strip()

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_docx(markdown_content):
    print("[2/4] Configurando estilos tipográficos (Verdana, UCM Blue) en python-docx...")
    doc = docx.Document()

    # Configuración de página (A4, márgenes 2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Encabezado
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Máster en Data Science & Big Data (UCM) | TFM AI-Dashboard Tenerife & TUI Group")
        hrun.font.name = "Verdana"
        hrun.font.size = Pt(8)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        # Pie de página
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Memoria Final de Trabajo de Fin de Máster — Curso 2025–2026")
        frun.font.name = "Verdana"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Colores corporativos
    COLOR_PRIMARY = RGBColor(31, 78, 121)     # #1F4E79 Azul Corporativo UCM
    COLOR_SECONDARY = RGBColor(47, 85, 151)   # #2F5597 Azul Medio
    COLOR_TEXT = RGBColor(38, 38, 38)         # #262626 Texto Oscuro
    COLOR_MUTED = RGBColor(89, 89, 89)        # #595959 Gris Neutro
    HEX_HEADER_BG = "1F4E79"
    HEX_ZEBRA_BG = "F2F4F7"
    HEX_CODE_BG = "F8F9FA"

    # Estilo Normal
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Verdana'
    font_normal.size = Pt(9.5)
    font_normal.color.rgb = COLOR_TEXT

    print("[3/4] Parseando Markdown y estructurando párrafos, tablas y bloques de código...")
    lines = markdown_content.split("\n")
    
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_code_block():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "\n".join(code_lines)
        # Añadir como párrafo formateado monospaced en caja
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        
        run = p.add_run(clean_xml_string(code_text))
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(33, 37, 41)
        code_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        # Parsear filas de tabla
        clean_rows = []
        for line in table_lines:
            if not line.strip() or line.strip().startswith("+--") or line.strip().startswith("+=="):
                continue
            if "|" in line:
                parts = [p.strip() for p in line.split("|")]
                # Si el primer y último elemento están vacíos debido a los bordes exteriores:
                if parts and parts[0] == "":
                    parts = parts[1:]
                if parts and parts[-1] == "":
                    parts = parts[:-1]
                # Comprobar si es fila separadora markdown (|---|---|)
                if all(re.match(r"^:?-+:?$", p) for p in parts if p):
                    continue
                if parts:
                    clean_rows.append(parts)
        
        if clean_rows:
            col_count = max(len(r) for r in clean_rows)
            # Rellenar filas más cortas
            for r in clean_rows:
                while len(r) < col_count:
                    r.append("")
            
            table = doc.add_table(rows=len(clean_rows), cols=col_count)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)

            for i, row_data in enumerate(clean_rows):
                row = table.rows[i]
                is_header = (i == 0)
                for j, cell_text in enumerate(row_data):
                    cell = row.cells[j]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    
                    # Formatear texto de celda
                    clean_text = clean_math_and_latex(cell_text.replace("**", "").replace("`", ""))
                    run = p.add_run(clean_xml_string(clean_text))
                    run.font.name = "Verdana"
                    
                    if is_header:
                        set_cell_background(cell, HEX_HEADER_BG)
                        run.font.bold = True
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        if i % 2 == 1:
                            set_cell_background(cell, HEX_ZEBRA_BG)
                        run.font.size = Pt(8)
                        run.font.color.rgb = COLOR_TEXT
                        # Alineación inteligente
                        if any(char.isdigit() for char in clean_text) and len(clean_text) < 15:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        else:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_lines = []

    # Iterar por cada línea de markdown
    for line in lines:
        stripped = line.strip()

        # Bloques de código
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code_block()
                continue
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Detección de tablas ASCII o Markdown
        if (stripped.startswith("|") and stripped.endswith("|")) or (stripped.startswith("+---") or stripped.startswith("+===")):
            if not in_table:
                in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            in_table = False
            flush_table()

        # Separador horizontal / Salto de página
        if stripped in ["---", "***", "___"]:
            doc.add_page_break()
            continue

        # Encabezados Markdown
        if stripped.startswith("# "):
            title_text = clean_math_and_latex(stripped[2:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml_string(title_text))
            run.font.name = "Verdana"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            continue

        if stripped.startswith("## "):
            title_text = clean_math_and_latex(stripped[3:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml_string(title_text))
            run.font.name = "Verdana"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY
            continue

        if stripped.startswith("### "):
            title_text = clean_math_and_latex(stripped[4:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml_string(title_text))
            run.font.name = "Verdana"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
            continue

        if stripped.startswith("#### "):
            title_text = clean_math_and_latex(stripped[5:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_xml_string(title_text))
            run.font.name = "Verdana"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = COLOR_MUTED
            continue

        # Listas con viñetas
        if stripped.startswith("* ") or stripped.startswith("- "):
            bullet_text = clean_math_and_latex(stripped[2:].strip())
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Formatear negritas dentro del bullet
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(clean_xml_string(part[2:-2]))
                    r.font.bold = True
                else:
                    r = p.add_run(clean_xml_string(part))
                r.font.name = "Verdana"
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_TEXT
            continue

        # Listas numeradas
        if re.match(r"^\d+\.\s", stripped):
            num_text = clean_math_and_latex(re.sub(r"^\d+\.\s", "", stripped))
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(clean_xml_string(part[2:-2]))
                    r.font.bold = True
                else:
                    r = p.add_run(clean_xml_string(part))
                r.font.name = "Verdana"
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_TEXT
            continue

        # Párrafos regulares de texto
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Descomponer enlaces, negritas y código en línea
            clean_line = clean_math_and_latex(re.sub(r'\[(.*?)\]\(.*?\)', r'\1', stripped)) # Simplificar links para Word
            parts = re.split(r'(\*\*.*?\*\*|\`.*?\`|\$.*?\$)', clean_line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(clean_xml_string(part[2:-2]))
                    r.font.bold = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(clean_xml_string(part[1:-1]))
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(180, 40, 40)
                else:
                    r = p.add_run(clean_xml_string(part))
                r.font.name = "Verdana"
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_TEXT

    # Asegurar flush de cualquier bloque remanente
    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    # Guardar documento
    output_docx = BASE_DIR / "TFM_Memoria_Final_UCM.docx"
    try:
        doc.save(str(output_docx))
        print(f"[4/4] Documento Word generado con éxito en: {output_docx}")
    except PermissionError:
        fallback_docx = BASE_DIR / "TFM_Memoria_Final_UCM_Actualizada.docx"
        doc.save(str(fallback_docx))
        print(f"[4/4] AVISO: '{output_docx.name}' está abierto en Microsoft Word.")
        print(f" -> Guardado exitosamente en archivo alternativo: {fallback_docx}")
        output_docx = fallback_docx
    return output_docx


if __name__ == "__main__":
    print("=" * 80)
    print("COMPILADOR MAESTRO DE LA MEMORIA DE TFM (UCM - MÁSTER BIG DATA & DATA SCIENCE)")
    print("=" * 80)
    md_content = assemble_markdown()
    docx_path = build_docx(md_content)
    print("=" * 80)
    print("PROCESO COMPLETADO EXITOSAMENTE.")
    print("Archivos listos para entrega académica y presentación ante TUI Group.")
    print("=" * 80)
