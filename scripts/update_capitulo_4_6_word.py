import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml

def update_docx():
    src_path = 'docs/capitulos_4_y_5_memoria_word_actualizado.docx'
    doc = docx.Document(src_path)

    # 1. Update Paragraph 21 (Intro to Section 4.6)
    p21_text = (
        "Conociendo los forzadores territoriales y el potencial PTNA, se articuló el pipeline oficial de "
        "clustering espacial en Python (analytics/clustering/run_hdbscan_clustering.py). La marcada orografía "
        "insular y la polarización turística invalidan algoritmos basados en particiones esféricas homogéneas "
        "(K-Means) o sin control legal de protección (que en modelos preliminares clasificaban erróneamente el Teide como urbano). "
        "El modelo definitivo opera sobre 8 covariables canónicas: plazas alojativas y radianza nocturna VIIRS "
        "(estabilizadas con transformación logarítmica log1p para atenuar colas pesadas), NDVI (vigor vegetal), "
        "NDBI (huella construida), altitud media, pendiente, distancia a la costa y el porcentaje de superficie en Espacio Natural Protegido (pct_area_enp, variable crítica). "
        "Tras estandarización con StandardScaler, un análisis PCA condensa el 81,5 % de la varianza en tres dimensiones no redundantes. "
        "La calibración óptima de HDBSCAN (min_cluster_size = 30, min_samples = 10, selección EOM) identificó de forma no supervisada "
        "4 macro-clústeres de densidad (57,3 % del territorio). El 42,7 % restante de celdas de transición y ruido se reasignó "
        "mediante reglas de experto territorial: celdas con ≥ 500 plazas pasaron a «Saturado / Overtourism»; celdas con VIIRS ≥ 20 nW/(cm²·sr) "
        "y ENP < 20 % se tipificaron como «Urbano Residencial»; y el ruido remanente se absorbió proyectándolo al centroide euclídeo más próximo "
        "en el espacio PCA. El resultado consolida seis tipologías territoriales exhaustivas con el 100 % de cobertura insular:"
    )
    
    # Check paragraph 21
    doc.paragraphs[21].text = p21_text
    # Re-apply formatting
    p21 = doc.paragraphs[21]
    p21.paragraph_format.space_before = Pt(0)
    p21.paragraph_format.space_after = Pt(6)
    p21.paragraph_format.line_spacing = 1.15
    p21.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p21.runs:
        run.font.name = 'Verdana'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(40, 40, 40)

    # 2. Update Table 3 (Table 0 in doc)
    # The new table data has 6 clusters + header
    t3_data = [
        ['Tipología Territorial', 'Cobertura y Peso (n / %)', 'Perfil Físico, Biofísico y Turístico', 'Comarcas y Municipios Representativos', 'Directriz Estratégica para TUI Group'],
        [
            'Espacio Natural / Teide y Cumbre',
            '716 celdas\n(27,8 %)',
            'Altitud media ~1.768 m, máxima protección (93 % ENP), radianza VIIRS nula (~0 nW), sustrato volcánico de alta cota.',
            'Parque Nacional del Teide, cumbre central y Corona Forestal alta.',
            'Preservación absoluta y exclusión hotelera; comercialización exclusiva de astroturismo y senderismo diurno de bajo impacto con guías certificados.'
        ],
        [
            'Espacios Rurales Protegidos (Anaga/Teno)',
            '456 celdas\n(17,7 %)',
            'Altitud media ~714 m, relieve muy abrupto (pendiente media ~30°), 87 % ENP, elevado vigor vegetal (NDVI alto), nula planta hotelera masiva.',
            'Macizo de Anaga (Santa Cruz, La Laguna), Parque Rural de Teno (Buenavista, Santiago del Teide).',
            'Ecoturismo botánico y senderismo regulado de residuo cero; valorización del patrimonio etnográfico sin nuevas edificaciones.'
        ],
        [
            'Transición Costera y Medianías',
            '895 celdas\n(34,7 %)',
            'Altitud media ~328 m, baja protección (8 % ENP), radianza VIIRS moderada (4,9 nW), posición bisagra entre el litoral y la cumbre.',
            'Corredores periurbanos y medianías bajas de Granadilla, San Miguel, Güímar, Fasnia y Candelaria.',
            'Vector de descompresión territorial; desarrollo de productos híbridos de media montaña, enoturismo y desvío de flujos excursionistas.'
        ],
        [
            'Rural Agrícola / Medianías Norte',
            '412 celdas\n(16,0 %)',
            'Altitud media ~490 m, alto verdor y humedad (elevado NDVI), baja edificación (NDBI bajo), VIIRS bajo (4,8 nW), clima templado con influencia atlántica.',
            'Medianías de Icod de los Vinos, La Orotava, Los Realejos, Garachico, Buenavista del Norte, Tacoronte.',
            'Foco prioritario de inversión y diversificación de TUI: micro-hoteles boutique rurales, agroturismo, rutas gastronómicas de guachinches y estancias de desconexión.'
        ],
        [
            'Saturado / Overtourism\n(Regla experto ≥500 pl.)',
            '61 celdas\n(2,4 %)',
            'Cota litoral (~71 m), hiperdensidad de camas (>500 pl/celda, >1.850 pl/km²), alta radianza VIIRS (~28 nW), queja dominante de masificación y ruido en NLP.',
            'Playa de las Américas, Los Cristianos, Costa Adeje (Adeje, Arona) y frente litoral maduro de Puerto de la Cruz.',
            'Moratoria y contención estricta de camas; reconversión de planta existente hacia estándares sostenibles (Travelife), elevación de ADR y redistribución de clientes.'
        ],
        [
            'Urbano Residencial\n(Regla experto VIIRS≥20)',
            '39 celdas\n(1,5 %)',
            'Máxima radianza económica (>41 nW promedio), 0 % ENP, elevada huella de suelo construido (NDBI), funciones residenciales y de servicios metropolitanos.',
            'Casco urbano consolidado de Santa Cruz de Tenerife y San Cristóbal de La Laguna.',
            'Turismo cultural urbano, patrimonio UNESCO, rutas gastronómicas metropolitanas y eventos corporativos MICE, protegiendo el parque residencial.'
        ]
    ]

    # Rebuild Table 0 in doc
    t0 = doc.tables[0]
    # Remove existing rows if needed or resize
    # To replace table cleanly, let's update cell contents or replace rows
    while len(t0.rows) < len(t3_data):
        t0.add_row()
    while len(t0.rows) > len(t3_data):
        # remove last row
        tr = t0.rows[-1]._tr
        tr.getparent().remove(tr)

    # Column widths for 5 columns
    t3_widths = [Inches(1.4), Inches(1.0), Inches(1.8), Inches(1.4), Inches(1.9)]
    
    # If table had 4 columns originally, let's adjust grid
    # Actually, replacing the table XML or inserting a newly formatted table right after P22 is cleaner:
    # Let's remove old table and insert new table
    tbl_element = t0._tbl
    parent = tbl_element.getparent()
    tbl_idx = parent.index(tbl_element)
    parent.remove(tbl_element)

    # Create new table in document at the exact position
    new_tbl = doc.add_table(rows=len(t3_data), cols=5)
    new_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(new_tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = t3_widths[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(t3_data[r_idx][c_idx])
            run.font.name = 'Verdana'
            if r_idx == 0:
                shd = parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="102C57"/>')
                tcPr.append(shd)
                run.font.bold = True
                run.font.size = Pt(8.0)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.size = Pt(7.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
                if r_idx % 2 == 1:
                    shd = parse_xml('<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F4F6F9"/>')
                    tcPr.append(shd)

    # Move new table element to the original table position
    new_tbl_elem = new_tbl._tbl
    parent.remove(new_tbl_elem)
    parent.insert(tbl_idx, new_tbl_elem)

    # 3. Add strategic narrative after table (between table 3 and heading 5)
    # Let's see where heading 5 is:
    h5_idx = None
    for i, p in enumerate(doc.paragraphs):
        if '5. Procesamiento del Lenguaje Natural' in p.text:
            h5_idx = i
            break
    
    # We can add a transitional paragraph explaining the 2D matrix and the 5 TUI product archetypes
    p_strat_text = (
        "Para transformar este diagnóstico territorial discreto en prescripciones accionables de negocio para TUI Group, "
        "los 2.579 hexágonos se proyectan en una matriz continua bidimensional: el Eje 1 (Saturación Turística [0–1], "
        "calibrado según plazas, VIIRS, proximidad litoral y densidad hotelera) frente al Eje 2 (Potencial Rural y Sostenible [0–1], "
        "que integra el índice PTNA derivado de MGWR, NDVI vegetal, ausencia de suelo sellado y gobernanza climática ESG). "
        "Este cruce analítico permite prescribir a cada microzona insular uno de los 5 Arquetipos de Producto Turístico de TUI "
        "(Sol y Playa Premium, Ecoturismo Rural y Medianías, Cultural y Patrimonial, Aventura y Turismo Activo, y Bienestar/Wellness), "
        "los cuales se explotan interactivamente en la vista de «Oportunidades TUI» del AI-Dashboard."
    )
    
    # Let's insert this paragraph before Heading 5
    if h5_idx is not None:
        p_target = doc.paragraphs[h5_idx]
        p_strat = p_target.insert_paragraph_before()
        p_strat.paragraph_format.space_before = Pt(6)
        p_strat.paragraph_format.space_after = Pt(8)
        p_strat.paragraph_format.line_spacing = 1.15
        p_strat.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_strat = p_strat.add_run(p_strat_text)
        r_strat.font.name = 'Verdana'
        r_strat.font.size = Pt(10)
        r_strat.font.color.rgb = RGBColor(40, 40, 40)

    # 4. Save to v2 and try to overwrite original
    v2_path = 'docs/capitulos_4_y_5_memoria_word_actualizado_v2.docx'
    doc.save(v2_path)
    print(f"Successfully saved {v2_path}")

    try:
        doc.save(src_path)
        print(f"Successfully overwritten {src_path}")
    except Exception as e:
        print(f"Notice: {src_path} is currently locked by Word ({e}). The updated document is available at {v2_path}")

if __name__ == '__main__':
    update_docx()
