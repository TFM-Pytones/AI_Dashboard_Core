import docx
from docx.shared import Pt
import win32com.client as win32
import os
import pdfplumber

src = "TFM_Memoria_Final_UCM.docx"
dst = "TFM_Memoria_Final_UCM_20caras.docx"
pdf = "TFM_Memoria_Final_UCM_20caras.pdf"

doc = docx.Document(src)

for s in doc.sections:
    s.top_margin = Pt(65)
    s.bottom_margin = Pt(65)
    s.left_margin = Pt(68)
    s.right_margin = Pt(68)

in_bib = False
for p in doc.paragraphs:
    txt = p.text.strip()
    if not txt:
        continue
    if any(k in txt for k in ["UNIVERSIDAD COMPLUTENSE", "Máster en Data Science", "TRABAJO DE FIN DE MÁSTER", "AI-Dashboard para la gestión", "Desafío 3 – Caso TUI Group", "Autores:", "Curso Académico"]):
        continue
    
    if txt == "Referencias bibliográficas":
        in_bib = True
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        continue
        
    if txt == "Anexos":
        in_bib = False
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        continue
        
    if in_bib:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.size = Pt(8.5)
    else:
        st = p.style.name if p.style else ""
        if st.startswith("Heading 1") or st.startswith("Título 1"):
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.12
        elif st.startswith("Heading 2") or st.startswith("Título 2"):
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.10
        elif st.startswith("Heading 3") or st.startswith("Título 3"):
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.10
        else:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4.0)
            p.paragraph_format.line_spacing = 1.18

# Tablas del cuerpo compactas
for t in doc.tables[:4]:
    t.autofit = True
    for row in t.rows:
        for cell in row.cells:
            for cp in cell.paragraphs:
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0.8)
                cp.paragraph_format.line_spacing = 1.0
                for cr in cp.runs:
                    if cr.font.size and cr.font.size > Pt(8.0):
                        cr.font.size = Pt(8.0)

doc.save(dst)

word = win32.gencache.EnsureDispatch('Word.Application')
word.Visible = False
word.DisplayAlerts = 0
try:
    wdoc = word.Documents.Open(os.path.abspath(dst), ReadOnly=True)
    wdoc.SaveAs2(os.path.abspath(pdf), 17)
    wdoc.Close(False)
finally:
    word.Quit()

with pdfplumber.open(pdf) as pdoc:
    print(f"Total pages: {len(pdoc.pages)}")
    res = {}
    for i in range(3, len(pdoc.pages)):
        text = pdoc.pages[i].extract_text() or ''
        for key in ['Resumen Ejecutivo', '1. Introducción y contexto', '8. Conclusiones', '8.3. Limitaciones', 'Referencias bibliográficas', 'Anexos']:
            if key.lower() in text.lower() and key not in res:
                res[key] = i + 1
    for k, v in res.items():
        print(f"   {k}: Página {v}")
    p_res = res.get('Resumen Ejecutivo')
    p_ref = res.get('Referencias bibliográficas')
    p_anx = res.get('Anexos')
    if p_res and p_ref:
        caras = p_ref - p_res + 1
        print(f"   -> CARAS DESDE RESUMEN HASTA BIBLIOGRAFÍA = {p_ref} - {p_res} + 1 = {caras} CARAS")
        print(f"   -> ANEXOS EN PÁGINA {p_anx}")
